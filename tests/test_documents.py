from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from alvis.core.models import Artifact
from alvis.extract import AutoExtractor, DocxExtractor, PdfExtractor, XlsxExtractor


def _artifact(content_type: str, data: bytes) -> Artifact:
    return Artifact(
        step_id="s1",
        uri="s3://alvis/report",
        content_type=content_type,
        data=data,
    )


@pytest.fixture()
def docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph("Revenue grew by 20%.")
    document.add_heading("Risks", level=2)
    document.add_paragraph("Currency exposure remains high.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revenue"
    sheet.append(["Year", "Amount"])
    sheet.append([2025, 1000])
    sheet.append([2026, 1200])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def pdf_bytes() -> bytes:
    return _minimal_pdf("Hello Alvis PDF")


def _minimal_pdf(text: str) -> bytes:
    """Build a minimal single-page PDF with one line of extractable text."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    obj = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        4: f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
        5: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    lines = [b"%PDF-1.4"]
    offsets: dict[int, int] = {}
    for number in sorted(obj):
        offsets[number] = len(b"\n".join(lines))
        lines.append(f"{number} 0 obj\n".encode() + obj[number] + b"\nendobj")
    xref_offset = len(b"\n".join(lines))
    xref = b"xref\n0 6\n0000000000 65535 f \n"
    xref += b"".join(
        f"{offsets[n]:010d} 00000 n \n".encode() for n in sorted(obj)
    )
    trailer = (
        b"\ntrailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref_offset).encode()
        + b"\n%%EOF"
    )
    return b"\n".join(lines) + b"\n" + xref + trailer


async def test_pdf_extractor(pdf_bytes: bytes) -> None:
    document = await PdfExtractor().extract(_artifact("application/pdf", pdf_bytes))
    assert document.sections[0].heading == "Page 1"
    assert "Hello Alvis" in document.sections[0].body


async def test_docx_extractor(docx_bytes: bytes) -> None:
    document = await DocxExtractor().extract(
        _artifact(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_bytes,
        )
    )
    headings = [s.heading for s in document.sections]
    assert headings == ["Executive Summary", "Risks"]
    assert any("Revenue grew by 20%" in s.body for s in document.sections)


async def test_xlsx_extractor(xlsx_bytes: bytes) -> None:
    document = await XlsxExtractor().extract(
        _artifact(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_bytes,
        )
    )
    assert document.sections[0].heading == "Revenue"
    assert "Year | Amount" in document.sections[0].body
    assert "2026 | 1200" in document.sections[0].body


async def test_auto_dispatches_by_content_type(
    docx_bytes: bytes, xlsx_bytes: bytes, pdf_bytes: bytes
) -> None:
    extractor = AutoExtractor()
    for content_type_url, data in (
        ("text/markdown", b"# Title\n\nbody\n"),
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_bytes,
        ),
        (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_bytes,
        ),
        ("application/pdf", pdf_bytes),
    ):
        document = await extractor.extract(_artifact(content_type_url, data))
        assert len(document.sections) >= 1


async def test_fs_source_ingests_documents(docx_bytes: bytes, tmp_path) -> None:
    from alvis.sources import FilesystemSource

    doc_path = tmp_path / "report.docx"
    doc_path.write_bytes(docx_bytes)
    source = FilesystemSource(path=str(tmp_path))
    artifacts = await source.fetch()
    assert len(artifacts) == 1
    assert artifacts[0].content_type.startswith(
        "application/vnd.openxmlformats-officedocument"
    )


async def test_e2e_pdf_pipeline(pdf_bytes: bytes, tmp_path) -> None:
    from alvis import dsl, run_async

    (tmp_path / "note.pdf").write_bytes(pdf_bytes)
    config = dsl.pipeline(dsl.fs(str(tmp_path)), index=dsl.memory())
    result = await run_async(config)
    assert result.documents_ingested == 1
    assert result.chunks_indexed >= 1